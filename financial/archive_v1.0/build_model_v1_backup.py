"""Builds DataMoat_Financial_Model.xlsx from assumptions.csv + seasonality.csv.

Every calculation is a live Excel formula driven by named ranges on the Assumptions sheet.
Edit the CSVs (or the blue cells in the workbook) — never the formulas.
Run:  python build_model.py
"""
import csv
import datetime
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "DataMoat_Financial_Model.xlsx")
VERSION = "R1-v1.0 · " + datetime.date.today().isoformat()

H = 60  # LTV horizon in months
SEGS = ["S", "M", "L"]
SEG_NAMES = {"S": "Small (<250)", "M": "Mid (250-800)", "L": "Large (>800)"}
SEG_VARS = {"Students", "Fee", "OnboardMult"}
TORNADO = ["ListPrice", "Fee", "DropoutRate", "UpfrontShare", "NInstalments", "InterventionSuccess",
           "Recall_Minor", "BlindMonths", "AdmissionUplift", "ITCShare", "ChurnInSeason", "RenewalSpike",
           "PilotConv", "PilotsPerRepMonth"]
CHURN_GRID = [0.01, 0.02, 0.03, 0.04, 0.05]
ACQ_GRID = [15000, 25000, 35000, 50000, 75000]
PRICE_GRID = [20, 30, 40, 50, 60, 80]
RED_GRID = [0.05, 0.08, 0.11, 0.15, 0.20, 0.25]

# ---------------------------------------------------------------- styles
ARIAL = "Arial"
F_IN = Font(name=ARIAL, size=10, color="0000FF")
F_CALC = Font(name=ARIAL, size=10, color="000000")
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_HDR = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=16, bold=True)
F_SUB = Font(name=ARIAL, size=10, italic=True, color="555555")
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="555555")
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

FMT = {
    "inr": '"₹"#,##0;("₹"#,##0);"-"',
    "inr2": '"₹"#,##0.00;("₹"#,##0.00);"-"',
    "pct": '0.0%;(0.0%);"-"',
    "int": '#,##0;(#,##0);"-"',
    "num1": '#,##0.0;(#,##0.0);"-"',
    "num2": '#,##0.00;(#,##0.00);"-"',
    "x": '0.00"x";(0.00"x");"-"',
    "mo": '0.0" mo";(0.0" mo");"-"',
    "txt": "@",
}


def put(ws, ref, value, font=F_CALC, fmt=None, fill=None, bold=False, align=None, border=False):
    c = ws[ref]
    c.value = value
    c.font = Font(name=font.name, size=font.size, color=font.color, bold=bold or font.bold,
                  italic=font.italic)
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = BOX
    return c


def header_row(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        col = chr(64 + start_col + i)
        put(ws, f"{col}{row}", lab, font=F_HDR, fill=FILL_HDR,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True))


def section(ws, row, text, last_col="K"):
    for col in range(1, ord(last_col) - 63):
        ws.cell(row=row, column=col).fill = FILL_SEC
    put(ws, f"A{row}", text, font=F_BOLD, fill=FILL_SEC)


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# ---------------------------------------------------------------- data
with open(os.path.join(HERE, "assumptions.csv"), encoding="utf-8-sig") as f:
    ASSUMP = list(csv.DictReader(f))
A_BY = {a["name"]: a for a in ASSUMP}
with open(os.path.join(HERE, "seasonality.csv"), encoding="utf-8-sig") as f:
    SEAS = list(csv.DictReader(f))


# ---------------------------------------------------------------- calculation specs
# kind: 'hdr' section label · 'in' input from Assumptions · 'link' computed value from another sheet
#       'calc' formula. f(R) returns the formula; R(key) resolves a key to a cell/name.
def I(name):
    a = A_BY[name + "_M"] if name in SEG_VARS else A_BY[name]
    label = a["label"].replace(" - Mid", "") if name in SEG_VARS else a["label"]
    return {"kind": "in", "key": name, "label": label, "unit": a["unit"], "fmt": a["fmt"]}


def LNK(key, label, unit, fmt):
    return {"kind": "link", "key": key, "label": label, "unit": unit, "fmt": fmt}


def C(key, label, unit, fmt, f, out=False):
    return {"kind": "calc", "key": key, "label": label, "unit": unit, "fmt": fmt, "f": f, "out": out}


def HD(text):
    return {"kind": "hdr", "key": "_h_" + text, "label": text}


TL0 = 37  # first timeline row on Seasonality


def tl(col, t):
    return f"Seasonality!${col}${TL0 + t - 1}"


def spec_cogs():
    return [
        HD("Inputs (from Assumptions)"),
        I("ReportsPerStudentMonth"), I("TokensInPerReport"), I("TokensOutPerReport"), I("LLMPriceIn"),
        I("LLMPriceOut"), I("WAMsgsPerStudentMonth"), I("WAPricePerMsg"), I("ComputePerStudentMonth"),
        I("ConsentCostPerStudentYear"), I("ServerCostMonth"), I("BSPFeeMonth"), I("InstPerStack"),
        I("CompliancePerInstMonth"), I("Students"),
        HD("Variable cost per active student per month"),
        C("cogs_llm", "LLM (AI parent reports, hosted model)", "₹/student/month", "inr2",
          lambda R: f"={R('ReportsPerStudentMonth')}*({R('TokensInPerReport')}*{R('LLMPriceIn')}"
                    f"+{R('TokensOutPerReport')}*{R('LLMPriceOut')})/1000000"),
        C("cogs_wa", "WhatsApp utility messages", "₹/student/month", "inr2",
          lambda R: f"={R('WAMsgsPerStudentMonth')}*{R('WAPricePerMsg')}"),
        C("cogs_compute", "Compute + storage + backups", "₹/student/month", "inr2",
          lambda R: f"={R('ComputePerStudentMonth')}"),
        C("cogs_consent", "Consent OTP / verification (annual ÷ 12)", "₹/student/month", "inr2",
          lambda R: f"={R('ConsentCostPerStudentYear')}/12"),
        C("COGS_PerStudent", "Total variable COGS per student-month", "₹/student/month", "inr2",
          lambda R: f"=SUM({R('cogs_llm')}:{R('cogs_consent')})", out=True),
        HD("Fixed cost per institute per month"),
        C("HostingPerInst", "Hosting share = (server + DB + WhatsApp BSP fee) ÷ institutes per stack",
          "₹/institute/month", "inr",
          lambda R: f"=({R('ServerCostMonth')}+{R('BSPFeeMonth')})/{R('InstPerStack')}"),
        C("COGS_FixedInst", "Hosting share + DPDP compliance", "₹/institute/month", "inr",
          lambda R: f"={R('HostingPerInst')}+{R('CompliancePerInstMonth')}", out=True),
        C("COGS_AllIn", "All-in COGS per student-month at this size", "₹/student/month", "inr2",
          lambda R: f"={R('COGS_PerStudent')}+{R('COGS_FixedInst')}/{R('Students')}", out=True),
    ]


def spec_pricing():
    return [
        HD("Inputs (from Assumptions)"),
        I("ListPrice"), I("MinFee"), I("NegotiationDisc"), I("PrepayUptake"), I("PrepayDisc"), I("GST"),
        HD("Realised price"),
        C("EffDisc", "Effective discount = negotiation + prepay uptake × prepay discount", "%", "pct",
          lambda R: f"={R('NegotiationDisc')}+{R('PrepayUptake')}*{R('PrepayDisc')}", out=True),
        C("RealisedPrice", "Realised price per student-month (ex-GST)", "₹/student/month", "inr2",
          lambda R: f"={R('ListPrice')}*(1-{R('EffDisc')})", out=True),
        C("RealisedPriceGST", "Price the owner feels (incl. GST)", "₹/student/month", "inr2",
          lambda R: f"={R('RealisedPrice')}*(1+{R('GST')})"),
        I("CompetitorAnchor"),
        C("PriceVsAnchor", "Our realised price ÷ what institutes already pay for an ERP/LMS", "x", "x",
          lambda R: f"=IF({R('CompetitorAnchor')}>0,{R('RealisedPrice')}/{R('CompetitorAnchor')},0)"),
    ]


def platform_fee_expr(R, price):
    s, mn = R("Students"), R("MinFee")
    return (f"SUMPRODUCT(({s}*Season_Idx*{price}>{mn})*({s}*Season_Idx*{price})"
            f"+({s}*Season_Idx*{price}<={mn})*{mn})*(1-{R('EffDisc')})")


def loss_vec(R):
    """Per-course-month fee lost if a student leaves that month (array expression over DropMonth).
    Instalment payers: unpaid instalments (N equal instalments, paid at the start of each 1/N of the course).
    Full payers: pro-rata refund of the remaining months, times the share of institutes that actually refund."""
    n = R("NInstalments")
    return (f"{R('Fee')}*((1-{R('UpfrontShare')})*(1-(INT((DropMonth-1)*{n}/CourseMonths)+1)/{n})"
            f"+{R('UpfrontShare')}*{R('RefundEnforced')}*DropRemain)")


def spec_roi():
    return [
        HD("Inputs (from Assumptions / Pricing)"),
        I("Students"), I("Fee"), I("DropoutRate"), I("UpfrontShare"), I("NInstalments"), I("RefundEnforced"),
        I("BlindMonths"), I("MinorShare"), I("Recall_Minor"), I("Recall_Adult"), I("Precision"),
        I("ActionRate"), I("InterventionSuccess"), I("AdmissionUplift"), I("MarginExtraAdmission"),
        I("StaffHoursSaved"), I("StaffHourCost"), I("InterventionMinutes"), I("MentorHourCost"),
        I("ListPrice"), I("MinFee"), I("GST"), I("ITCShare"),
        LNK("EffDisc", "Effective discount (Pricing)", "%", "pct"),
        HD("1. What a dropout actually costs (unpaid instalments + refunds to full payers)"),
        C("Dropouts", "Dropouts per year", "students/yr", "num1",
          lambda R: f"={R('Students')}*{R('DropoutRate')}"),
        C("FeeAtRisk", "Avg fee lost per dropout (timing curve × unpaid instalments / pro-rata refund)",
          "₹/dropout", "inr", lambda R: f"=SUMPRODUCT(DropTiming,{loss_vec(R)})", out=True),
        C("GrossLoss", "Fee lost to dropouts per year (no platform)", "₹/yr", "inr",
          lambda R: f"={R('Dropouts')}*{R('FeeAtRisk')}"),
        HD("2. Prevention funnel (replaces the guide's flat 25%)"),
        C("Recall", "Blended recall (minors: academic model + rules; 18+: full ML)", "%", "pct",
          lambda R: f"={R('MinorShare')}*{R('Recall_Minor')}+(1-{R('MinorShare')})*{R('Recall_Adult')}"),
        C("CatchableShare", "Share of dropouts after the 'blind' months (enough test data to flag)", "%", "pct",
          lambda R: f"=SUMPRODUCT(DropTiming*(DropMonth>{R('BlindMonths')}))"),
        C("EffReduction", "Effective dropout reduction = recall × action × RCT effect × catchable", "%", "pct",
          lambda R: f"={R('Recall')}*{R('ActionRate')}*{R('InterventionSuccess')}*{R('CatchableShare')}", out=True),
        C("Prevented", "Dropouts prevented per year", "students/yr", "num1",
          lambda R: f"={R('Dropouts')}*{R('EffReduction')}", out=True),
        C("FeeSavedPerPrevented", "Fee saved per prevented dropout (catchable months only)", "₹/student", "inr",
          lambda R: f"=IF({R('CatchableShare')}>0,SUMPRODUCT(DropTiming*(DropMonth>{R('BlindMonths')})"
                    f"*{loss_vec(R)})/{R('CatchableShare')},0)", out=True),
        HD("3. Value to the institute (₹/yr)"),
        C("FeesRetained", "Fees retained", "₹/yr", "inr",
          lambda R: f"={R('Prevented')}*{R('FeeSavedPerPrevented')}"),
        C("AdmissionValue", "Margin on extra admissions (results / parent trust)", "₹/yr", "inr",
          lambda R: f"={R('Students')}*{R('AdmissionUplift')}*{R('Fee')}*{R('MarginExtraAdmission')}"),
        C("StaffValue", "Staff time saved", "₹/yr", "inr",
          lambda R: f"={R('StaffHoursSaved')}*{R('StaffHourCost')}*12"),
        C("TotalValue", "Total value", "₹/yr", "inr",
          lambda R: f"={R('FeesRetained')}+{R('AdmissionValue')}+{R('StaffValue')}", out=True),
        HD("4. Cost to the institute (₹/yr)"),
        C("PlatformFee", "Platform fee (seasonal active students, min fee, discounts; ex-GST)", "₹/yr", "inr",
          lambda R: "=" + platform_fee_expr(R, R("ListPrice"))),
        C("GSTCost", "GST the institute cannot reclaim", "₹/yr", "inr",
          lambda R: f"={R('PlatformFee')}*{R('GST')}*(1-{R('ITCShare')})"),
        C("Flagged", "Students flagged per year (true + false alarms)", "students/yr", "num1",
          lambda R: f"=IF({R('Precision')}>0,{R('Dropouts')}*{R('Recall')}/{R('Precision')},0)"),
        C("InterventionCost", "Mentor time spent acting on flags", "₹/yr", "inr",
          lambda R: f"={R('Flagged')}*{R('ActionRate')}*{R('InterventionMinutes')}/60*{R('MentorHourCost')}"),
        C("TotalCost", "Total cost", "₹/yr", "inr",
          lambda R: f"={R('PlatformFee')}+{R('GSTCost')}+{R('InterventionCost')}", out=True),
        HD("5. Result"),
        C("NetBenefit", "Net benefit to institute", "₹/yr", "inr",
          lambda R: f"={R('TotalValue')}-{R('TotalCost')}", out=True),
        C("ROI", "Institute ROI (value ÷ cost)", "x", "x",
          lambda R: f"=IF({R('TotalCost')}>0,{R('TotalValue')}/{R('TotalCost')},0)", out=True),
        C("FeesOnlyROI", "Fees-only ROI (no soft benefits)", "x", "x",
          lambda R: f"=IF({R('TotalCost')}>0,{R('FeesRetained')}/{R('TotalCost')},0)", out=True),
        C("BreakEvenStudents", "Break-even: students to save per year (fees only)", "students/yr", "num1",
          lambda R: f"=IF({R('FeeSavedPerPrevented')}>0,{R('TotalCost')}/{R('FeeSavedPerPrevented')},0)",
          out=True),
        C("BreakEvenReduction", "Break-even dropout reduction needed", "%", "pct",
          lambda R: f"=IF({R('Dropouts')}>0,{R('BreakEvenStudents')}/{R('Dropouts')},0)", out=True),
        C("UpliftNeeded", "Admissions uplift needed to break even, after fees + staff time", "% of students",
          "pct", lambda R: f"=IF({R('Students')}*{R('Fee')}*{R('MarginExtraAdmission')}>0,MAX(0,{R('TotalCost')}"
                           f"-{R('FeesRetained')}-{R('StaffValue')})/({R('Students')}*{R('Fee')}*{R('MarginExtraAdmission')}),0)",
          out=True),
        C("UpliftNeededStudents", "…in extra admissions per year", "students/yr", "num1",
          lambda R: f"={R('UpliftNeeded')}*{R('Students')}", out=True),
        C("CostPerStudentYr", "Platform cost per student per year (incl. unreclaimed GST)", "₹/student/yr", "inr",
          lambda R: f"=({R('PlatformFee')}+{R('GSTCost')})/{R('Students')}", out=True),
        C("CostPctFee", "Platform cost as % of the student's fee", "%", "pct",
          lambda R: f"={R('CostPerStudentYr')}/{R('Fee')}", out=True),
    ]


def spec_ue():
    items = [
        HD("Inputs (from Assumptions / Pricing / COGS_Build)"),
        I("Students"), I("ListPrice"), I("MinFee"), I("BadDebt"), I("GatewayFee"),
        I("ChurnInSeason"), I("PrepayUptake"), I("RenewalSpike"), I("CSM_CTC"), I("InstPerCSM"), I("RepCTC"),
        I("RepTravel"), I("PilotsPerRepMonth"), I("PilotConv"), I("PilotMonths"), I("PilotSetupCost"),
        I("OnboardBase"),
        I("OnboardMult"), I("MarketingPerPaid"),
        LNK("EffDisc", "Effective discount (Pricing)", "%", "pct"),
        LNK("COGS_PerStudent", "Variable COGS per student-month (COGS_Build)", "₹/student/month", "inr2"),
        LNK("COGS_FixedInst", "Fixed COGS per institute-month (COGS_Build)", "₹/month", "inr"),
        HD("1. Year-1 economics per institute (before churn)"),
        C("Y1Revenue", "Collected revenue, first 12 paid months (ex-GST, after bad debt)", "₹/yr", "inr",
          lambda R: f"=SUM({R('rev_1')}:{R('rev_12')})", out=True),
        C("AvgMonthlyRev", "Average monthly revenue", "₹/month", "inr",
          lambda R: f"={R('Y1Revenue')}/12", out=True),
        C("Y1COGS", "COGS, first 12 months (students × seasonality × COGS + fixed + gateway)", "₹/yr", "inr",
          lambda R: f"={R('Students')}*{R('COGS_PerStudent')}*SUM(Seasonality!$C${TL0}:$C${TL0 + 11})"
                    f"+12*{R('COGS_FixedInst')}+{R('GatewayFee')}*{R('Y1Revenue')}"),
        C("GrossMargin", "Gross margin", "%", "pct",
          lambda R: f"=IF({R('Y1Revenue')}>0,({R('Y1Revenue')}-{R('Y1COGS')})/{R('Y1Revenue')},0)", out=True),
        C("SupportPerInst", "Customer-success cost per institute", "₹/month", "inr",
          lambda R: f"={R('CSM_CTC')}/{R('InstPerCSM')}"),
        C("AvgMonthlyContr", "Average monthly contribution (after COGS + support)", "₹/month", "inr",
          lambda R: f"=SUM({R('contr_1')}:{R('contr_12')})/12", out=True),
        C("ContrMargin", "Contribution margin", "%", "pct",
          lambda R: f"=IF({R('AvgMonthlyRev')}>0,{R('AvgMonthlyContr')}/{R('AvgMonthlyRev')},0)", out=True),
        HD("2. Cost to acquire one paying institute"),
        C("SalesPerPilot", "Sales cost per pilot started (rep CTC + travel ÷ pilots)", "₹", "inr",
          lambda R: f"=({R('RepCTC')}+{R('RepTravel')})/{R('PilotsPerRepMonth')}"),
        C("OnboardCost", "Onboarding cost (data migration + training)", "₹", "inr",
          lambda R: f"={R('OnboardBase')}*{R('OnboardMult')}"),
        C("PilotCOGS", "Cost of serving the free pilot", "₹", "inr",
          lambda R: f"={R('PilotMonths')}*({R('Students')}*{R('COGS_PerStudent')}+{R('COGS_FixedInst')})"),
        C("CAC", "CAC: sales + marketing per paid institute (failed pilots included)", "₹", "inr",
          lambda R: f"={R('SalesPerPilot')}/{R('PilotConv')}+{R('MarketingPerPaid')}", out=True),
        C("AcqCost", "Fully loaded acquisition cost = (sales + pilot setup + pilot COGS) ÷ conversion "
                     "+ onboarding (paid only) + marketing", "₹", "inr",
          lambda R: f"=({R('SalesPerPilot')}+{R('PilotSetupCost')}+{R('PilotCOGS')})/{R('PilotConv')}"
                    f"+{R('OnboardCost')}+{R('MarketingPerPaid')}", out=True),
        HD("3. Retention (seasonal, not 1/churn)"),
        C("ChurnEff", "In-season churn, monthly payers only (prepaid annual plans cannot leave mid-year)",
          "%/month", "pct", lambda R: f"={R('ChurnInSeason')}*(1-{R('PrepayUptake')})"),
        C("SpikeMonthly", "Extra churn per renewal month", "%/month", "pct",
          lambda R: f"=1-(1-{R('RenewalSpike')})^(1/NRenew)"),
        C("AnnualRetention", "Annual logo retention", "%", "pct",
          lambda R: f"=(1-{R('ChurnEff')})^12*(1-{R('RenewalSpike')})", out=True),
        C("ExpLifetime", "Expected paid lifetime within horizon", "months", "num1",
          lambda R: f"=SUM({R('surv_1')}:{R(f'surv_{H}')})", out=True),
        HD(f"4. Lifetime value ({H}-month horizon, survival-weighted, discounted)"),
        C("LTV", "LTV (contribution)", "₹", "inr", lambda R: f"={R(f'cum_{H}')}", out=True),
        C("LTVtoCAC", "LTV ÷ fully loaded acquisition cost", "x", "x",
          lambda R: f"=IF({R('AcqCost')}>0,{R('LTV')}/{R('AcqCost')},0)", out=True),
        C("Payback", "Payback (months of paid life, discounted)", "months", "int",
          lambda R: f'=IF({R(f"cum_{H}")}<{R("AcqCost")},{H + 1},COUNTIF({R("cum_1")}:{R(f"cum_{H}")},"<"&{R("AcqCost")})+1)',
          out=True),
        HD("5. Month-by-month engine (t = paid month; calendar from Seasonality timeline)"),
    ]
    for t in range(1, H + 1):
        items.append(C(f"rev_{t}", f"Revenue collected · month {t}", "₹", "inr",
                       lambda R, t=t: f"=MAX({R('MinFee')},{R('Students')}*{tl('C', t)}*{R('ListPrice')})"
                                      f"*(1-{R('EffDisc')})*(1-{R('BadDebt')})"))
    for t in range(1, H + 1):
        items.append(C(f"contr_{t}", f"Contribution · month {t}", "₹", "inr",
                       lambda R, t=t: f"={R(f'rev_{t}')}-({R('Students')}*{tl('C', t)}*{R('COGS_PerStudent')}"
                                      f"+{R('COGS_FixedInst')}+{R('GatewayFee')}*{R(f'rev_{t}')})-{R('SupportPerInst')}"))
    for t in range(1, H + 1):
        if t == 1:
            f = lambda R: "=1"
        else:
            f = (lambda R, t=t: f"={R(f'surv_{t - 1}')}*(1-{R('ChurnEff')})"
                                f"*(1-{tl('D', t - 1)}*{R('SpikeMonthly')})")
        items.append(C(f"surv_{t}", f"Share of institutes still paying · month {t}", "%", "pct", f))
    for t in range(1, H + 1):
        prev = "" if t == 1 else None
        items.append(C(f"cum_{t}", f"Cumulative discounted contribution · month {t}", "₹", "inr",
                       lambda R, t=t: ("=" if t == 1 else f"={R(f'cum_{t - 1}')}+")
                                      + f"{R(f'surv_{t}')}*{R(f'contr_{t}')}*{tl('E', t)}"))
    return items


# ---------------------------------------------------------------- emitter
def emit(ws, items, start_row, cols, getters, write_labels=True, skip_inputs=False, extra_rows=None,
         label_col="B", unit_col="C", input_font=F_LINK):
    """Write items down the sheet; one column per getter. Returns key -> row."""
    rows, r = {}, start_row
    for it in items:
        if skip_inputs and it["kind"] in ("in", "link"):
            continue
        rows[it["key"]] = r
        r += 1
    allrows = dict(extra_rows or {})
    allrows.update(rows)
    for col, get in zip(cols, getters):
        def R(k, col=col, get=get):
            if k in allrows:
                return f"{col}{allrows[k]}"
            return get(k)
        for it in items:
            if it["key"] not in rows or (skip_inputs and it["kind"] in ("in", "link")):
                continue
            row = rows[it["key"]]
            if it["kind"] == "hdr":
                if write_labels:
                    for cc in range(1, 4 + len(cols)):
                        ws.cell(row=row, column=cc).fill = FILL_SEC
                    put(ws, f"{label_col}{row}", it["label"], font=F_BOLD, fill=FILL_SEC)
                continue
            if write_labels:
                put(ws, f"{label_col}{row}", it["label"])
                put(ws, f"{unit_col}{row}", it["unit"], font=F_NOTE)
            if it["kind"] in ("in", "link"):
                put(ws, f"{col}{row}", "=" + get(it["key"]), font=input_font, fmt=it["fmt"])
            else:
                put(ws, f"{col}{row}", it["f"](R), fmt=it["fmt"],
                    fill=FILL_OUT if it.get("out") else None)
    return rows


def seg_getter(seg):
    return lambda k: f"{k}_{seg}" if k in SEG_VARS else k


# ================================================================ workbook
wb = Workbook()
wb.remove(wb.active)
order = ["Cover", "Assumptions", "Seasonality", "Pricing", "COGS_Build", "Institute_ROI", "Unit_Econ",
         "Scenarios", "Sensitivity", "Checks", "Sources", "Engine"]
WS = {n: wb.create_sheet(n) for n in order}

# ---------------------------------------------------------------- Cover (scenario cell first: it is referenced everywhere)
cv = WS["Cover"]
add_name(wb, "Scenario", "Cover!$C$6")

# ---------------------------------------------------------------- Assumptions
asn = WS["Assumptions"]
put(asn, "A1", "Assumptions — every input in the model lives here", font=F_TITLE)
put(asn, "A2", "Blue = input you may edit (Bear / Base / Bull). Live = CHOOSE(Scenario, Bear, Base, Bull). "
               "Bear = worst case for OUR business. Type: Researched/Public = sourced · Guide = case brief · "
               "Assumption = unverified · Structural = modelling convention.", font=F_SUB)
header_row(asn, 4, ["Name (named range → Live)", "Input", "Unit", "Bear", "Base", "Bull", "LIVE", "Type",
                    "Confidence", "Source", "Note", "Pct out of [0,1]", "Live outside Bear–Bull"])
A_ROW = {}
r = 5
cur_sec = None
for a in ASSUMP:
    if a["section"] != cur_sec:
        cur_sec = a["section"]
        section(asn, r, cur_sec, last_col="M")
        r += 1
    A_ROW[a["name"]] = r
    put(asn, f"A{r}", a["name"], font=F_NOTE)
    put(asn, f"B{r}", a["label"])
    put(asn, f"C{r}", a["unit"], font=F_NOTE)
    for col, key in zip("DEF", ["bear", "base", "bull"]):
        put(asn, f"{col}{r}", float(a[key]), font=F_IN, fmt=a["fmt"])
    put(asn, f"G{r}", f"=CHOOSE(Scenario,D{r},E{r},F{r})", fmt=a["fmt"], fill=FILL_KEY, bold=True)
    put(asn, f"H{r}", a["type"], font=F_NOTE)
    put(asn, f"I{r}", a["confidence"], font=F_NOTE)
    put(asn, f"J{r}", a["source"], font=F_NOTE)
    put(asn, f"K{r}", a["note"], font=F_NOTE)
    put(asn, f"L{r}", f"=IF(OR(G{r}<0,G{r}>1),1,0)" if a["fmt"] == "pct" else 0, font=F_NOTE)
    put(asn, f"M{r}", f"=IF(AND(G{r}>=MIN(D{r},F{r}),G{r}<=MAX(D{r},F{r})),0,1)", font=F_NOTE)
    add_name(wb, a["name"], f"Assumptions!$G${r}")
    r += 1
A_LAST = r - 1
for col, w in zip("ABCDEFGHIJKLM", [24, 52, 16, 11, 11, 11, 12, 12, 10, 60, 40, 9, 9]):
    asn.column_dimensions[col].width = w
asn.freeze_panes = "D5"

# ---------------------------------------------------------------- Seasonality
se = WS["Seasonality"]
put(se, "A1", "Seasonality, dropout timing and the 60-month timeline", font=F_TITLE)
put(se, "A2", "Session month 1 = April. Blue cells are inputs (currently assumptions, see Source).", font=F_SUB)
put(se, "A4", "A. Active-student index and renewal (churn) months", font=F_BOLD)
header_row(se, 4, ["No.", "Month", "Active-student index", "Renewal month (1/0)", "Source"], start_col=1)
put(se, "A3", "A. Active-student index and renewal (churn) months", font=F_BOLD)
idx_rows = [s for s in SEAS if s["kind"] == "student_idx"]
flag_rows = {s["index"]: s for s in SEAS if s["kind"] == "renewal_flag"}
for i, s in enumerate(idx_rows):
    rr = 5 + i
    put(se, f"A{rr}", int(s["index"]))
    put(se, f"B{rr}", s["label"])
    put(se, f"C{rr}", float(s["value"]), font=F_IN, fmt="num2")
    put(se, f"D{rr}", int(flag_rows[s["index"]]["value"]), font=F_IN, fmt="0")
    put(se, f"E{rr}", s["source"] + (" | flag: " + flag_rows[s["index"]]["source"]
                                      if flag_rows[s["index"]]["source"] else ""), font=F_NOTE)
add_name(wb, "Season_Idx", "Seasonality!$C$5:$C$16")
add_name(wb, "Season_Flag", "Seasonality!$D$5:$D$16")
put(se, "B18", "Renewal months per year")
put(se, "C18", "=SUM(Season_Flag)", fmt="0", fill=FILL_OUT)
add_name(wb, "NRenew", "Seasonality!$C$18")

put(se, "A20", "B. When in the course students drop out", font=F_BOLD)
header_row(se, 21, ["Course month", "", "Share of dropouts", "Share of course left", "Source"], start_col=1)
drops = [s for s in SEAS if s["kind"] == "drop_timing"]
for i, s in enumerate(drops):
    rr = 22 + i
    put(se, f"A{rr}", int(s["index"]))
    put(se, f"B{rr}", s["label"], font=F_NOTE)
    put(se, f"C{rr}", float(s["value"]), font=F_IN, fmt="pct")
    put(se, f"D{rr}", f"=MAX(0,(CourseMonths-A{rr})/CourseMonths)", fmt="pct")
    put(se, f"E{rr}", s["source"], font=F_NOTE)
add_name(wb, "DropTiming", "Seasonality!$C$22:$C$31")
add_name(wb, "DropMonth", "Seasonality!$A$22:$A$31")
add_name(wb, "DropRemain", "Seasonality!$D$22:$D$31")
put(se, "B32", "Total (must be 100%)", font=F_BOLD)
put(se, "C32", "=SUM(DropTiming)", fmt="pct", fill=FILL_OUT)

put(se, "A34", "C. Timeline for LTV: paid month t → calendar month, student index, renewal flag, discount factor",
    font=F_BOLD)
header_row(se, 36, ["t (paid month)", "Calendar month #", "Active-student index", "Renewal flag",
                    "Discount factor"], start_col=1)
for t in range(1, H + 1):
    rr = TL0 + t - 1
    put(se, f"A{rr}", t)
    put(se, f"B{rr}", f"=MOD(StartMonth+A{rr}-2,12)+1")
    put(se, f"C{rr}", f"=INDEX(Season_Idx,B{rr})", fmt="num2")
    put(se, f"D{rr}", f"=INDEX(Season_Flag,B{rr})", fmt="0")
    put(se, f"E{rr}", f"=1/(1+DiscountRate)^(A{rr}/12)", fmt="0.000")
for col, w in zip("ABCDE", [14, 18, 20, 18, 70]):
    se.column_dimensions[col].width = w

# ---------------------------------------------------------------- Pricing
pr = WS["Pricing"]
put(pr, "A1", "Pricing — list price, discounts, realised price", font=F_TITLE)
put(pr, "A2", "Green = linked input · black = formula · green fill = output used elsewhere", font=F_SUB)
header_row(pr, 4, ["", "Item", "Unit", "LIVE"])
pr_rows = emit(pr, spec_pricing(), 5, ["D"], [seg_getter("M")])
for k in ("EffDisc", "RealisedPrice"):
    add_name(wb, k, f"Pricing!$D${pr_rows[k]}")
r0 = max(pr_rows.values()) + 3
put(pr, f"A{r0}", "Packaging (for the deck; Pro price is the model's list price)", font=F_BOLD)
header_row(pr, r0 + 1, ["Tier", "Contents", "Price", "Role"])
tiers = [
    ("Starter", "CSV / OMR upload, test-level reports, 1 batch", 0, "Free — lead generation, data on-ramp"),
    ("Pro (core)", "Topic heatmap, at-risk list with reasons, template parent reports", "=ListPrice",
     "Core paid plan, ₹/active student/month, min fee applies"),
    ("Premium (AI)", "GenAI Hindi/English parent reports on WhatsApp, percentile band, study plans",
     "Pro + add-on", "Upsell — NOT in base revenue until priced by research"),
    ("Annual plan", "Pro billed yearly upfront", "=ListPrice*(1-PrepayDisc)",
     "Cash up front, lower churn at renewal"),
]
for i, (a_, b_, c_, d_) in enumerate(tiers):
    rr = r0 + 2 + i
    put(pr, f"A{rr}", a_, bold=True)
    put(pr, f"B{rr}", b_)
    put(pr, f"C{rr}", c_, fmt="inr")
    put(pr, f"D{rr}", d_, font=F_NOTE)
for col, w in zip("ABCD", [16, 62, 20, 50]):
    pr.column_dimensions[col].width = w

# ---------------------------------------------------------------- COGS_Build
cg = WS["COGS_Build"]
put(cg, "A1", "COGS build-up — cost to serve one student (must match TECH_PLAN §16 / NB09)", font=F_TITLE)
put(cg, "A2", "Mid-segment size used for the all-in figure. Tech plan target band: ₹3–10 per student-month.",
    font=F_SUB)
header_row(cg, 4, ["", "Item", "Unit", "LIVE"])
cg_rows = emit(cg, spec_cogs(), 5, ["D"], [seg_getter("M")])
for k in ("COGS_PerStudent", "COGS_FixedInst", "COGS_AllIn"):
    add_name(wb, k, f"COGS_Build!$D${cg_rows[k]}")
for col, w in zip("ABCD", [4, 58, 20, 14]):
    cg.column_dimensions[col].width = w

# ---------------------------------------------------------------- Institute_ROI
ir = WS["Institute_ROI"]
put(ir, "A1", "Institute ROI — does it pay for the owner? (per segment, live scenario)", font=F_TITLE)
put(ir, "A2", "Same price and same dropout funnel as Unit_Econ. Green fill = key outputs.", font=F_SUB)
header_row(ir, 4, ["", "Line", "Unit"] + [SEG_NAMES[s] for s in SEGS])
ir_rows = emit(ir, spec_roi(), 5, ["D", "E", "F"], [seg_getter(s) for s in SEGS])
g0 = max(ir_rows.values()) + 3
put(ir, f"A{g0}", "The guide's naive model (for contrast — all inputs from the Guide block on Assumptions)",
    font=F_BOLD)
guide_ir = [
    ("gi_drop", "Dropouts per year", "=G_Students*G_Dropout", "num1"),
    ("gi_prev", "Dropouts prevented (flat 25%)", "=G_Students*G_Dropout*G_Prevent", "num1"),
    ("gi_saved", "Fees saved", "=G_Students*G_Dropout*G_Prevent*G_RemainingFee", "inr"),
    ("gi_cost", "Platform cost (12 months, no GST, no seasonality)", "=G_Students*G_Price*G_Months", "inr"),
    ("gi_roi", "Guide ROI", "=IF(ROUND(G_Students*G_Price*G_Months,6)>0,G_Students*G_Dropout*G_Prevent*G_RemainingFee/(G_Students*G_Price*G_Months),0)", "x"),
    ("gi_be", "Guide break-even students", "=G_Students*G_Price*G_Months/G_RemainingFee", "num1"),
]
GI = {}
for i, (k, lab, f, fm) in enumerate(guide_ir):
    rr = g0 + 1 + i
    GI[k] = f"Institute_ROI!$E${rr}"
    put(ir, f"B{rr}", lab)
    put(ir, f"E{rr}", f, fmt=fm)
for col, w in zip("ABCDEF", [3, 70, 14, 16, 16, 16]):
    ir.column_dimensions[col].width = w
ir.freeze_panes = "D5"

# ---------------------------------------------------------------- Unit_Econ
ue = WS["Unit_Econ"]
put(ue, "A1", "Unit economics per institute (live scenario)", font=F_TITLE)
put(ue, "A2", "LTV uses a seasonal survival curve (monthly churn + renewal spike), discounted, capped at "
              f"{H} months. Payback = {H + 1} means 'not within {H} months'.", font=F_SUB)
header_row(ue, 4, ["", "Line", "Unit"] + [SEG_NAMES[s] for s in SEGS])
ue_rows = emit(ue, spec_ue(), 5, ["D", "E", "F"], [seg_getter(s) for s in SEGS])
first_block = ue_rows["rev_1"]
ue.row_dimensions.group(first_block, ue_rows[f"cum_{H}"], hidden=True, outline_level=1)
# guide + break-even scale block to the right
put(ue, "H4", "Guide's naive unit economics", font=F_HDR, fill=FILL_HDR)
put(ue, "I4", "Value", font=F_HDR, fill=FILL_HDR)
guide_ue = [
    ("Revenue per institute per month", "=G_Students*G_Price", "inr"),
    ("Contribution per month", "=G_Students*G_Price-G_Students*G_Infra-G_Support", "inr"),
    ("CAC + onboarding", "=G_CAC+G_Onboard", "inr"),
    ("Payback (months)", "=(G_CAC+G_Onboard)/(G_Students*G_Price-G_Students*G_Infra-G_Support)", "num1"),
    ("LTV = contribution ÷ churn", "=(G_Students*G_Price-G_Students*G_Infra-G_Support)/G_Churn", "inr"),
    ("LTV ÷ (CAC + onboarding)", "=((G_Students*G_Price-G_Students*G_Infra-G_Support)/G_Churn)/(G_CAC+G_Onboard)", "x"),
]
GU = {}
for i, (lab, f, fm) in enumerate(guide_ue):
    rr = 5 + i
    GU[lab] = f"Unit_Econ!$I${rr}"
    put(ue, f"H{rr}", lab)
    put(ue, f"I{rr}", f, fmt=fm)
put(ue, "H13", "Break-even scale (fixed team cost ÷ Mid contribution)", font=F_HDR, fill=FILL_HDR)
put(ue, "I13", "Value", font=F_HDR, fill=FILL_HDR)
be = [
    ("Fixed team + tools cost per month", "=Founders*FounderStipend+Engineers*EngineerCTC+ToolsOpex", "inr"),
    ("Mid-institute avg monthly contribution", f"=E{ue_rows['AvgMonthlyContr']}", "inr"),
    ("Institutes needed to cover fixed cost", f"=IF(I15>0,I14/I15,0)", "num1"),
    ("…equivalent active students", "=I16*Students_M", "int"),
    ("Paid institutes one rep adds per month", "=PilotsPerRepMonth*PilotConv", "num1"),
]
for i, (lab, f, fm) in enumerate(be):
    rr = 14 + i
    put(ue, f"H{rr}", lab)
    put(ue, f"I{rr}", f, fmt=fm, fill=FILL_OUT if i == 2 else None)
put(ue, "H20", "Full hiring plan, cohorts, cash and funding need: Round 2 sheets.", font=F_NOTE)
for col, w in zip("ABCDEFGHI", [3, 66, 16, 16, 16, 16, 3, 50, 14]):
    ue.column_dimensions[col].width = w
ue.freeze_panes = "D5"

# ---------------------------------------------------------------- Engine (scenario / sensitivity re-runs)
en = WS["Engine"]
put(en, "A1", "Engine — re-runs the full chain for scenarios, tornado and churn grid. Do not edit.", font=F_TITLE)
eng_items = [it for it in spec_cogs() + spec_pricing() + spec_roi() + spec_ue() if it["kind"] != "hdr"]
eng_inputs = []
for it in eng_items:
    if it["kind"] == "in" and it["key"] not in eng_inputs:
        eng_inputs.append(it["key"])
cases = []
for s in SEGS:
    cases.append({"label": f"Live {s}", "seg": s, "scen": "=Scenario", "ov": "", "ovv": ""})
for sc, scn in [(1, "Bear"), (2, "Base"), (3, "Bull")]:
    for s in SEGS:
        cases.append({"label": f"{scn} {s}", "seg": s, "scen": sc, "ov": "", "ovv": ""})
def mid_name(nm):
    return nm + "_M" if nm in SEG_VARS else nm


for nm in TORNADO:
    rr = A_ROW[mid_name(nm)]
    cases.append({"label": f"T {nm} bear", "seg": "M", "scen": 2, "ov": nm, "ovv": f"=Assumptions!$D${rr}"})
    cases.append({"label": f"T {nm} bull", "seg": "M", "scen": 2, "ov": nm, "ovv": f"=Assumptions!$F${rr}"})
for j, _ in enumerate(CHURN_GRID):
    cases.append({"label": f"Churn grid {j + 1}", "seg": "M", "scen": "=Scenario", "ov": "ChurnInSeason",
                  "ovv": None, "grid": j})
for lbl, rr in [("Case", 1), ("Segment", 2), ("Scenario (1-3)", 3), ("Override input", 4), ("Override value", 5)]:
    put(en, f"B{rr + 1}", lbl, font=F_BOLD)


def colname(i):
    s = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


EIN0 = 8
ein_rows = {nm: EIN0 + i for i, nm in enumerate(eng_inputs)}
CASE_COL = {}
for ci, cs in enumerate(cases):
    col = colname(3 + ci)  # from D
    CASE_COL[cs["label"]] = col
    put(en, f"{col}2", cs["label"], font=F_BOLD)
    put(en, f"{col}3", cs["seg"])
    put(en, f"{col}4", cs["scen"])
    put(en, f"{col}5", cs["ov"])
    if cs.get("grid") is not None:
        cs["ovv"] = f"=Sensitivity!{chr(66 + cs['grid'])}$SENSROW$"
    put(en, f"{col}6", cs["ovv"])
    for nm in eng_inputs:
        src = f"{nm}_{cs['seg']}" if nm in SEG_VARS else nm
        ar = A_ROW[src]
        put(en, f"{col}{ein_rows[nm]}",
            f'=IF(${col}$5="{nm}",${col}$6,INDEX(Assumptions!$D${ar}:$F${ar},${col}$4))',
            fmt=A_BY[src]["fmt"])
for nm in eng_inputs:
    put(en, f"B{ein_rows[nm]}", nm, font=F_NOTE)
calc_start = EIN0 + len(eng_inputs) + 1
eng_rows = emit(en, eng_items, calc_start, [CASE_COL[c["label"]] for c in cases],
                [lambda k: (_ for _ in ()).throw(KeyError(k))] * len(cases),
                skip_inputs=True, extra_rows=ein_rows, input_font=F_CALC)
en.column_dimensions["B"].width = 40
en.freeze_panes = "D7"


def E(case, key):
    return f"Engine!${CASE_COL[case]}${eng_rows[key] if key in eng_rows else ein_rows[key]}"


# ---------------------------------------------------------------- Scenarios
sc = WS["Scenarios"]
put(sc, "A1", "Scenarios — honest model vs the guide's naive model", font=F_TITLE)
put(sc, "A2", "Bear / Base / Bull are always shown side by side, whatever the Cover selector says. "
              "Mid institute = 400 students (the guide's example).", font=F_SUB)
header_row(sc, 4, ["Metric (Mid institute)", "Unit", "Guide (naive)", "Bear", "Base", "Bull"])
metrics = [
    ("Effective dropout reduction", "%", "EffReduction", "=Assumptions!$E$" + str(A_ROW["G_Prevent"]), "pct"),
    ("Dropouts prevented per year", "students", "Prevented", "=" + GI["gi_prev"], "num1"),
    ("Fee lost per dropout", "₹", "FeeAtRisk", "=G_RemainingFee", "inr"),
    ("Fees retained per year", "₹/yr", "FeesRetained", "=" + GI["gi_saved"], "inr"),
    ("Total value to institute", "₹/yr", "TotalValue", "=" + GI["gi_saved"], "inr"),
    ("Total cost to institute (fee + GST + mentor time)", "₹/yr", "TotalCost", "=" + GI["gi_cost"], "inr"),
    ("Institute ROI", "x", "ROI", "=" + GI["gi_roi"], "x"),
    ("Fees-only ROI", "x", "FeesOnlyROI", "=" + GI["gi_roi"], "x"),
    ("Break-even students saved per year", "students", "BreakEvenStudents", "=" + GI["gi_be"], "num1"),
    ("Avg monthly revenue per institute", "₹/month", "AvgMonthlyRev", "=" + GU["Revenue per institute per month"], "inr"),
    ("Avg monthly contribution", "₹/month", "AvgMonthlyContr", "=" + GU["Contribution per month"], "inr"),
    ("Fully loaded acquisition cost", "₹", "AcqCost", "=" + GU["CAC + onboarding"], "inr"),
    ("Annual logo retention", "%", "AnnualRetention", "=(1-G_Churn)^12", "pct"),
    ("LTV", "₹", "LTV", "=" + GU["LTV = contribution ÷ churn"], "inr"),
    ("LTV ÷ acquisition cost", "x", "LTVtoCAC", "=" + GU["LTV ÷ (CAC + onboarding)"], "x"),
    ("Payback", "months", "Payback", "=" + GU["Payback (months)"], "num1"),
]
SC_ROW = {}
for i, (lab, unit, key, gf, fm) in enumerate(metrics):
    rr = 5 + i
    SC_ROW[key] = rr
    put(sc, f"A{rr}", lab)
    put(sc, f"B{rr}", unit, font=F_NOTE)
    put(sc, f"C{rr}", gf, fmt=fm, font=F_LINK)
    for col, scn in zip("DEF", ["Bear", "Base", "Bull"]):
        put(sc, f"{col}{rr}", "=" + E(f"{scn} M", key), fmt=fm, font=F_LINK,
            fill=FILL_OUT if key in ("ROI", "LTVtoCAC", "BreakEvenStudents") else None)
r0 = 5 + len(metrics) + 2
put(sc, f"A{r0}", "By segment", font=F_BOLD)
for block, (key, fm, title) in enumerate([("ROI", "x", "Institute ROI"), ("LTVtoCAC", "x", "LTV ÷ acquisition cost"),
                                          ("Payback", "num1", "Payback (months)")]):
    rb = r0 + 1 + block * 6
    header_row(sc, rb, [title, "", "", "Bear", "Base", "Bull"])
    for i, s in enumerate(SEGS):
        rr = rb + 1 + i
        put(sc, f"A{rr}", SEG_NAMES[s])
        for col, scn in zip("DEF", ["Bear", "Base", "Bull"]):
            put(sc, f"{col}{rr}", "=" + E(f"{scn} {s}", key), fmt=fm, font=F_LINK)
for col, w in zip("ABCDEF", [52, 12, 16, 14, 14, 14]):
    sc.column_dimensions[col].width = w

# ---------------------------------------------------------------- Sensitivity
sn = WS["Sensitivity"]
put(sn, "A1", "Sensitivity — what actually moves the answer", font=F_TITLE)
put(sn, "A2", "Tornado: one input at a time moved to its Bear / Bull value, all else Base (Mid institute). "
              "Grids follow the live scenario.", font=F_SUB)
header_row(sn, 4, ["Input", "Bear value", "Bull value", "ROI @ Bear", "ROI @ Bull", "ROI swing",
                   "LTV/CAC @ Bear", "LTV/CAC @ Bull", "LTV/CAC swing"])
put(sn, "A5", "BASE (all inputs Base)", bold=True)
put(sn, "D5", "=" + E("Base M", "ROI"), fmt="x", font=F_LINK)
put(sn, "E5", "=" + E("Base M", "ROI"), fmt="x", font=F_LINK)
put(sn, "G5", "=" + E("Base M", "LTVtoCAC"), fmt="x", font=F_LINK)
put(sn, "H5", "=" + E("Base M", "LTVtoCAC"), fmt="x", font=F_LINK)
for i, nm in enumerate(TORNADO):
    rr = 6 + i
    ar = A_ROW[mid_name(nm)]
    fm = A_BY[mid_name(nm)]["fmt"]
    put(sn, f"A{rr}", A_BY[mid_name(nm)]["label"])
    put(sn, f"B{rr}", f"=Assumptions!$D${ar}", fmt=fm, font=F_LINK)
    put(sn, f"C{rr}", f"=Assumptions!$F${ar}", fmt=fm, font=F_LINK)
    put(sn, f"D{rr}", "=" + E(f"T {nm} bear", "ROI"), fmt="x", font=F_LINK)
    put(sn, f"E{rr}", "=" + E(f"T {nm} bull", "ROI"), fmt="x", font=F_LINK)
    put(sn, f"F{rr}", f"=ABS(E{rr}-D{rr})", fmt="x")
    put(sn, f"G{rr}", "=" + E(f"T {nm} bear", "LTVtoCAC"), fmt="x", font=F_LINK)
    put(sn, f"H{rr}", "=" + E(f"T {nm} bull", "LTVtoCAC"), fmt="x", font=F_LINK)
    put(sn, f"I{rr}", f"=ABS(H{rr}-G{rr})", fmt="x")
T_LAST = 6 + len(TORNADO) - 1
for col, w in zip("ABCDEFGHI", [52, 12, 12, 12, 12, 12, 14, 14, 14]):
    sn.column_dimensions[col].width = w

ch = BarChart()
ch.type = "bar"
ch.title = "Institute ROI: Bear vs Bull value of each input (Mid, rest Base)"
ch.y_axis.title = "ROI (x)"
ch.add_data(Reference(sn, min_col=4, max_col=5, min_row=4, max_row=T_LAST), titles_from_data=True)
ch.set_categories(Reference(sn, min_col=1, min_row=6, max_row=T_LAST))
ch.height, ch.width = 9, 22
sn.add_chart(ch, "K4")

g1 = T_LAST + 3
put(sn, f"A{g1}", "Grid 1 — Institute ROI (Mid, live scenario): list price × effective dropout reduction", font=F_BOLD)
put(sn, f"A{g1 + 1}", "Rows: effective reduction · Columns: list price ₹/student/month (blue = grid points)",
    font=F_NOTE)
mid = {k: f"Institute_ROI!$E${ir_rows[k]}" for k in ir_rows}
for j, p in enumerate(PRICE_GRID):
    put(sn, f"{chr(66 + j)}{g1 + 2}", p, font=F_IN, fmt="inr", fill=FILL_SEC)
for i, red in enumerate(RED_GRID):
    rr = g1 + 3 + i
    put(sn, f"A{rr}", red, font=F_IN, fmt="pct", fill=FILL_SEC)
    for j in range(len(PRICE_GRID)):
        col = chr(66 + j)
        price = f"{col}${g1 + 2}"
        R_mid = lambda k: mid[k] if k in mid else (f"{k}_M" if k in SEG_VARS else k)
        fee = platform_fee_expr(R_mid, price)
        put(sn, f"{col}{rr}",
            f"=({mid['Dropouts']}*{mid['FeeSavedPerPrevented']}*$A{rr}+{mid['AdmissionValue']}+{mid['StaffValue']})"
            f"/({fee}*(1+GST*(1-ITCShare))+{mid['InterventionCost']})", fmt="x")
G1_RANGE = f"B{g1 + 3}:{chr(65 + len(PRICE_GRID))}{g1 + 2 + len(RED_GRID)}"

g2 = g1 + 3 + len(RED_GRID) + 2
put(sn, f"A{g2}", "Grid 2 — LTV ÷ acquisition cost (Mid, live scenario): monthly churn × acquisition cost",
    font=F_BOLD)
put(sn, f"A{g2 + 1}", "Monthly in-season churn (blue) → LTV from Engine", font=F_NOTE)
SENSROW = g2 + 1
for j, c_ in enumerate(CHURN_GRID):
    col = chr(66 + j)
    put(sn, f"{col}{g2 + 1}", c_, font=F_IN, fmt="pct", fill=FILL_SEC)
    put(sn, f"{col}{g2 + 2}", "=" + E(f"Churn grid {j + 1}", "LTV"), fmt="inr", font=F_LINK)
put(sn, f"A{g2 + 2}", "LTV at that churn (₹)", font=F_NOTE)
put(sn, f"A{g2 + 3}", "Acquisition cost ₹ ↓", font=F_NOTE)
for i, acq in enumerate(ACQ_GRID):
    rr = g2 + 4 + i
    put(sn, f"A{rr}", acq, font=F_IN, fmt="inr", fill=FILL_SEC)
    for j in range(len(CHURN_GRID)):
        col = chr(66 + j)
        put(sn, f"{col}{rr}", f"={col}${g2 + 2}/$A{rr}", fmt="x")
G2_RANGE = f"B{g2 + 4}:{chr(65 + len(CHURN_GRID))}{g2 + 3 + len(ACQ_GRID)}"
# patch engine override links to the churn grid header row
for ci, cs in enumerate(cases):
    if cs.get("grid") is not None:
        cell = en[f"{CASE_COL[cs['label']]}6"]
        cell.value = cell.value.replace("$SENSROW$", f"${SENSROW}")
put(sn, f"A{g2 + 10}", "Reading: ROI < 1.0x red, 1–2x amber, ≥ 2x green. LTV/CAC < 3x red.", font=F_NOTE)

red_f = PatternFill("solid", fgColor="F8CBAD")
amb_f = PatternFill("solid", fgColor="FFE699")
grn_f = PatternFill("solid", fgColor="C6EFCE")
for rng in [G1_RANGE, "D5:E" + str(T_LAST)]:
    sn.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["1"], fill=red_f))
    sn.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["1", "1.9999"], fill=amb_f))
    sn.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=grn_f))
for rng in [G2_RANGE, "G5:H" + str(T_LAST)]:
    sn.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["3"], fill=red_f))
    sn.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["3"], fill=grn_f))

# ---------------------------------------------------------------- Checks
ck = WS["Checks"]
put(ck, "A1", "Checks — any red flag means the model is not safe to quote", font=F_TITLE)
header_row(ck, 3, ["Check", "Type", "Value", "Flag (1 = fail)"])
mid_ue = lambda k: f"Unit_Econ!$E${ue_rows[k]}"
checks = [
    ("Scenario selector is 1, 2 or 3", "Error", "=Scenario", "=IF(OR(Scenario=1,Scenario=2,Scenario=3),0,1)"),
    ("Dropout timing curve sums to 100%", "Error", "=SUM(DropTiming)", "=IF(ABS(C5-1)<0.0001,0,1)"),
    ("Percent inputs within 0–100%", "Error", f"=SUM(Assumptions!L5:L{A_LAST})", "=IF(C6=0,0,1)"),
    ("Live values within Bear–Bull range", "Error", f"=SUM(Assumptions!M5:M{A_LAST})", "=IF(C7=0,0,1)"),
    ("At least one renewal month (avoids ÷0)", "Error", "=NRenew", "=IF(C8>=1,0,1)"),
    ("Active-student index between 0 and 1.2", "Error",
     "=SUMPRODUCT((Season_Idx<0)+(Season_Idx>1.2))", "=IF(C9=0,0,1)"),
    ("Effective reduction ≤ RCT effect (Mid)", "Error",
     f"=Institute_ROI!E{ir_rows['EffReduction']}-InterventionSuccess", "=IF(C10<=0.000001,0,1)"),
    ("Instalments ≥ 1 and ≤ course months", "Error", "=NInstalments",
     "=IF(AND(C{r}>=1,C{r}<=CourseMonths,C{r}=INT(C{r})),0,1)"),
]
for s in SEGS:
    col = "DEF"[SEGS.index(s)]
    checks.append((f"Engine reproduces Institute_ROI ROI ({s})", "Error",
                   f"=ABS(Institute_ROI!{col}{ir_rows['ROI']}-{E('Live ' + s, 'ROI')})",
                   None))
    checks.append((f"Engine reproduces Unit_Econ LTV ({s})", "Error",
                   f"=ABS(Unit_Econ!{col}{ue_rows['LTV']}-{E('Live ' + s, 'LTV')})", None))
checks += [
    ("Survival curve stays within 0–1 (Mid)", "Error",
     f"=IF(AND(MIN(Unit_Econ!E{ue_rows['surv_1']}:E{ue_rows[f'surv_{H}']})>=0,"
     f"MAX(Unit_Econ!E{ue_rows['surv_1']}:E{ue_rows[f'surv_{H}']})<=1),0,1)", "=C{r}"),
    ("Guide ROI reproduces the brief's 2.25x", "Error", "=" + GI["gi_roi"], "=IF(ABS(C{r}-2.25)<0.01,0,1)"),
    ("WARNING: RCT effect exceeds share of dropouts that are preventable", "Warning",
     "=InterventionSuccess-Preventable", "=IF(C{r}<=0,0,1)"),
    ("WARNING: list price above the highest ERP/LMS anchor found (₹21/student/month)", "Warning",
     "=ListPrice", "=IF(C{r}<=CompetitorAnchorHigh,0,1)"),
    ("WARNING: months with negative contribution (Mid)", "Warning",
     f"=COUNTIF(Unit_Econ!E{ue_rows['contr_1']}:E{ue_rows[f'contr_{H}']},\"<0\")", "=IF(C{r}=0,0,1)"),
    ("WARNING: all-in COGS outside tech band ₹3–10 per student-month", "Warning", "=COGS_AllIn",
     "=IF(AND(C{r}>=3,C{r}<=10),0,1)"),
    ("WARNING: Mid institute ROI below 1.0x", "Warning", f"=Institute_ROI!E{ir_rows['ROI']}", "=IF(C{r}>=1,0,1)"),
    ("WARNING: Mid LTV/CAC below 3x", "Warning", f"=Unit_Econ!E{ue_rows['LTVtoCAC']}", "=IF(C{r}>=3,0,1)"),
    ("WARNING: Mid payback above 24 months", "Warning", f"=Unit_Econ!E{ue_rows['Payback']}", "=IF(C{r}<=24,0,1)"),
]
for i, (lab, typ, val, flag) in enumerate(checks):
    rr = 4 + i
    put(ck, f"A{rr}", lab)
    put(ck, f"B{rr}", typ, font=F_NOTE)
    put(ck, f"C{rr}", val.replace("{r}", str(rr)), fmt="0.000")
    if flag is None:
        flag = "=IF(C{r}<0.000001,0,1)"
    put(ck, f"D{rr}", flag.replace("{r}", str(rr)), fmt="0", bold=True)
CK_LAST = 4 + len(checks) - 1
put(ck, f"A{CK_LAST + 2}", "ERRORS", bold=True)
put(ck, f"D{CK_LAST + 2}", f'=SUMIF(B4:B{CK_LAST},"Error",D4:D{CK_LAST})', fmt="0", bold=True, fill=FILL_KEY)
put(ck, f"A{CK_LAST + 3}", "WARNINGS (honest red flags to discuss, not bugs)", bold=True)
put(ck, f"D{CK_LAST + 3}", f'=SUMIF(B4:B{CK_LAST},"Warning",D4:D{CK_LAST})', fmt="0", bold=True)
add_name(wb, "ErrorCount", f"Checks!$D${CK_LAST + 2}")
add_name(wb, "WarningCount", f"Checks!$D${CK_LAST + 3}")
ck.conditional_formatting.add(f"D4:D{CK_LAST + 3}",
                              CellIsRule(operator="greaterThan", formula=["0"], fill=red_f))
for col, w in zip("ABCD", [64, 10, 14, 14]):
    ck.column_dimensions[col].width = w

# ---------------------------------------------------------------- Sources
so = WS["Sources"]
put(so, "A1", "Sources & evidence register", font=F_TITLE)
put(so, "A2", "Raw research is in Financial_Model/research/*.md. Anything typed 'Assumption' is unverified — "
              "say so on the slide.", font=F_SUB)
header_row(so, 4, ["Name", "Input", "Type", "Confidence", "Source", "Note"])
for i, a in enumerate(ASSUMP):
    rr = 5 + i
    for col, k in zip("ABCDEF", ["name", "label", "type", "confidence", "source", "note"]):
        put(so, f"{col}{rr}", a[k], font=F_NOTE if col in "AF" else F_CALC)
rr = 5 + len(ASSUMP) + 1
put(so, f"A{rr}", "Evidence mix", bold=True)
for j, t in enumerate(["Researched", "Public", "Guide", "Assumption", "Structural"]):
    put(so, f"B{rr + 1 + j}", t)
    put(so, f"C{rr + 1 + j}", f'=COUNTIF(C5:C{4 + len(ASSUMP)},"{t}")', fmt="0")
for col, w in zip("ABCDEF", [24, 50, 12, 11, 80, 40]):
    so.column_dimensions[col].width = w
so.freeze_panes = "A5"

# ---------------------------------------------------------------- Cover (content)
put(cv, "A1", "The Data Moat — Financial Model", font=F_TITLE)
put(cv, "A2", f"Institute Intelligence Platform · Nextute × Finance Club IIT Patna, Case B · {VERSION}", font=F_SUB)
put(cv, "B6", "SCENARIO  (1 = Bear, 2 = Base, 3 = Bull)  →", bold=True)
put(cv, "C6", 2, font=F_IN, fill=FILL_KEY, fmt="0", border=True)
put(cv, "D6", '=CHOOSE(Scenario,"Bear","Base","Bull")', bold=True)
dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
cv.add_data_validation(dv)
dv.add("C6")
put(cv, "B7", "Model checks", bold=True)
put(cv, "C7", '=IF(ErrorCount=0,"ALL CHECKS PASS","ERRORS: "&ErrorCount)', bold=True)
put(cv, "D7", '="Warnings to discuss: "&WarningCount', font=F_NOTE)
cv.conditional_formatting.add("C7", CellIsRule(operator="equal", formula=['"ALL CHECKS PASS"'], fill=grn_f))
header_row(cv, 9, ["", "Headline (Mid institute, 400 students, live scenario)", "Our model", "Guide (naive)"])
heads = [
    ("Effective dropout reduction", f"=Institute_ROI!E{ir_rows['EffReduction']}", "=G_Prevent", "pct"),
    ("Dropouts prevented per year", f"=Institute_ROI!E{ir_rows['Prevented']}", "=" + GI["gi_prev"], "num1"),
    ("Institute ROI (all benefits)", f"=Institute_ROI!E{ir_rows['ROI']}", "=" + GI["gi_roi"], "x"),
    ("Institute ROI (fees only)", f"=Institute_ROI!E{ir_rows['FeesOnlyROI']}", "=" + GI["gi_roi"], "x"),
    ("Break-even: students to save per year", f"=Institute_ROI!E{ir_rows['BreakEvenStudents']}",
     "=" + GI["gi_be"], "num1"),
    ("Extra admissions needed to break even (after fees + staff time)",
     f"=Institute_ROI!E{ir_rows['UpliftNeededStudents']}", "n/a", "num1"),
    ("Platform cost as % of a student's fee", f"=Institute_ROI!E{ir_rows['CostPctFee']}",
     "=G_Price*G_Months/Fee_M", "pct"),
    ("Avg monthly revenue per institute", f"=Unit_Econ!E{ue_rows['AvgMonthlyRev']}",
     "=" + GU["Revenue per institute per month"], "inr"),
    ("Contribution margin", f"=Unit_Econ!E{ue_rows['ContrMargin']}",
     "=" + GU["Contribution per month"] + "/" + GU["Revenue per institute per month"], "pct"),
    ("Fully loaded acquisition cost", f"=Unit_Econ!E{ue_rows['AcqCost']}", "=" + GU["CAC + onboarding"], "inr"),
    ("LTV", f"=Unit_Econ!E{ue_rows['LTV']}", "=" + GU["LTV = contribution ÷ churn"], "inr"),
    ("LTV ÷ acquisition cost", f"=Unit_Econ!E{ue_rows['LTVtoCAC']}", "=" + GU["LTV ÷ (CAC + onboarding)"], "x"),
    ("Payback (months)", f"=Unit_Econ!E{ue_rows['Payback']}", "=" + GU["Payback (months)"], "num1"),
    ("Institutes needed to cover fixed team cost", "=Unit_Econ!I16", "n/a", "num1"),
]
for i, (lab, f, g, fm) in enumerate(heads):
    rr = 10 + i
    put(cv, f"B{rr}", lab)
    put(cv, f"C{rr}", f, fmt=fm, font=F_LINK, fill=FILL_OUT, border=True)
    put(cv, f"D{rr}", g, fmt=fm, font=F_LINK, border=True)
r0 = 10 + len(heads) + 1
legend = [
    ("Colour code", ""),
    ("Blue text", "Input — edit only these (Assumptions D:F, Seasonality, grid points, Scenario)"),
    ("Black text", "Formula"),
    ("Green text", "Link from another sheet"),
    ("Yellow fill", "Key lever / live value"),
    ("Green fill", "Key output"),
    ("", ""),
    ("How to use", ""),
    ("1", "Pick a scenario in C6. Everything recalculates; Scenarios and Sensitivity always show all three."),
    ("2", "Change an input on Assumptions (Bear/Base/Bull) — never type over a formula."),
    ("3", "Cover C7 must say ALL CHECKS PASS before you quote any number."),
    ("4", "To rebuild from the CSVs: python build_model.py (Financial_Model folder)."),
    ("", ""),
    ("Evidence labels", "Researched/Public = sourced (see Sources) · Guide = case brief · "
                        "Assumption = unverified · Structural = modelling convention"),
    ("Scope", "Round 1: Pricing, COGS, Institute ROI, Unit economics, Scenarios, Sensitivity. "
              "Round 2 adds cohorts, hiring plan, 3-year P&L, cash and funding need."),
]
for i, (a_, b_) in enumerate(legend):
    rr = r0 + i
    put(cv, f"B{rr}", a_, bold=a_ in ("Colour code", "How to use", "Evidence labels", "Scope"))
    put(cv, f"C{rr}", b_, font=F_NOTE)
cv[f"B{r0 + 1}"].font = F_IN
cv[f"B{r0 + 3}"].font = F_LINK
cv[f"B{r0 + 4}"].fill = FILL_KEY
cv[f"B{r0 + 5}"].fill = FILL_OUT
for col, w in zip("ABCD", [3, 48, 22, 22]):
    cv.column_dimensions[col].width = w
cv["C6"].comment = Comment("Scenario selector: 1 Bear, 2 Base, 3 Bull", "model")

WS["Engine"].sheet_state = "hidden"
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
print("engine cases:", len(cases), "assumption rows:", len(ASSUMP))
