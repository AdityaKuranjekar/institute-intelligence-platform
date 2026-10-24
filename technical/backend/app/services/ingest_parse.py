"""File parsing for uploads: CSV and XLSX of any width, messy headers (English, Hinglish, Devanagari), merged
title rows, formulas, totals rows, hidden rows, text numbers. Pure functions, no database.

Reuses the helpers and synonym map of datagen/messy/reference_parse.py. Row semantics and error codes follow
that parser so that data/dirty/expected/*.json stays the golden reference (tested in tests/test_ingest_golden.py).
"""
from __future__ import annotations

import csv
import datetime as dt
import difflib
import io
import re
import sys
from typing import Any

import openpyxl

from .. import config

if str(config.ROOT) not in sys.path:
    sys.path.insert(0, str(config.ROOT))
from datagen.messy import reference_parse as ref  # noqa: E402

REV, ABS_SET, PRES_SET = ref._build_reverse_synonym_map()
EXTRA_SYNONYMS = {
    "response": ["response", "answer", "given answer", "marked", "option marked", "student answer", "उत्तर"],
    "question_no": ["q no", "qno", "question number", "प्रश्न क्रमांक"],
    "marks_awarded": ["marks awarded"],
}
for canon, syns in EXTRA_SYNONYMS.items():
    for s in syns:
        REV.setdefault(s.lower(), canon)
Q_COL = re.compile(r"^\s*(?:q|ques|question|qn|प्र|प्रश्न)[\s._\-]*(\d{1,3})\s*$", re.IGNORECASE)
BLANK_MARKERS = {"", "-", "--", "na", "n/a", "blank", "छोड़ा"}

REQUIRED = {
    "students": ["roll_no", "name", "dob", "join_date", "batch_name"],
    "guardians": ["roll_no", "guardian_name", "phone"],
    "attendance": ["roll_no", "attendance_date", "present_flag"],
    "fees": ["roll_no", "instalment_no", "due_date", "amount_paise"],
    "results_marks_only": ["test_name", "roll_no"],
    "results_responses": ["test_name", "roll_no"],
    "results_long": ["test_name", "roll_no", "question_no", "response"],
    "answer_key": ["question_no", "accepted_answers"],
}
KIND_TO_TYPES = {
    "students": {"students"}, "guardians": {"guardians"}, "attendance": {"attendance"}, "fees": {"fees"},
    "results": {"results_marks_only", "results_responses", "results_long"}, "answer-key": {"answer_key"},
}
DB_TYPE = {"results_long": "results_responses"}


class ParseError(ValueError):
    pass


def parse_date(val: Any) -> str | None:
    if isinstance(val, dt.datetime):
        return val.date().isoformat()
    if isinstance(val, dt.date):
        return val.isoformat()
    if val is None:
        return None
    s = str(val).strip()
    r = ref._parse_date_flexible(s)
    if r:
        return r
    for fmt in ("%d.%m.%Y", "%d %b %Y", "%d-%b-%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _num(val: Any, as_int: bool = False):
    s = str(val).replace(",", "").strip()
    f = float(s)
    return int(f) if as_int else f


# ---------------------------------------------------------------- reading
def read_rows(data: bytes, filename: str) -> list[tuple[int, list[Any], bool]]:
    """(1-based row number, values, hidden) for the sheet or CSV that holds the table."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        try:
            wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            wb_v = None
        best, best_score = None, -1
        for sh in wb_f.sheetnames:
            ws = wb_f[sh]
            score = 0
            for r_idx in range(1, min(12, ws.max_row + 1)):
                vals = [ws.cell(r_idx, c).value for c in range(1, ws.max_column + 1)]
                score = max(score, sum(1 for v in vals if v is not None and (str(v).strip().lower() in REV or Q_COL.match(str(v)))))
            if score > best_score:
                best, best_score = sh, score
        ws, wv = wb_f[best], (wb_v[best] if wb_v else None)
        rows = []
        for r_idx in range(1, ws.max_row + 1):
            vals = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r_idx, c).value
                if isinstance(v, str) and v.startswith("=") and wv is not None and wv.cell(r_idx, c).value is not None:
                    v = wv.cell(r_idx, c).value  # cached result of a formula when Excel saved one
                vals.append(v)
            rows.append((r_idx, vals, bool(ws.row_dimensions[r_idx].hidden)))
        return rows
    if name.endswith((".csv", ".txt", ".tsv")):
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:4096]
        delim = ","
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            pass
        return [(i + 1, row, False) for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim))]
    raise ParseError("unsupported file type: use .csv or .xlsx")


# ---------------------------------------------------------------- header mapping
def map_header(cell: Any, override: dict[str, str] | None = None) -> tuple[str | None, float, str]:
    """(canonical, confidence, method) for one header cell."""
    if cell is None or str(cell).strip() == "":
        return None, 0.0, "empty"
    raw = str(cell).strip()
    key = raw.lower()
    if override and raw in override:
        return override[raw], 1.0, "manual"
    if key in REV:
        return REV[key], 1.0, "synonym"
    m = Q_COL.match(raw)
    if m:
        return f"q{int(m.group(1))}", 0.95, "question_column"
    close = difflib.get_close_matches(key, list(REV), n=1, cutoff=0.82)
    if close:
        return REV[close[0]], round(difflib.SequenceMatcher(None, key, close[0]).ratio() * 0.9, 2), "fuzzy"
    return None, 0.0, "unmapped"


def find_header(rows, override=None, min_hits: int = 3):
    for pos, (_r, vals, _h) in enumerate(rows):
        mapping = {}
        for c_i, v in enumerate(vals):
            canon, conf, method = map_header(v, override)
            if canon and method != "fuzzy" or (canon and conf >= 0.75):
                mapping[c_i] = (canon, conf, method, str(v).strip())
        hits = len({m[0] for m in mapping.values()})
        if hits >= min_hits or (override and hits >= 2):
            return pos, mapping
    raise ParseError("could not find a header row with at least 3 recognisable columns; pick the columns manually")


def detect_type(canon: set[str]) -> str:
    if "dob" in canon and "join_date" in canon:
        return "students"
    if "guardian_name" in canon:
        return "guardians"
    if "attendance_date" in canon:
        return "attendance"
    if "instalment_no" in canon:
        return "fees"
    if "question_no" in canon and ("accepted_answers" in canon or "q_type" in canon):
        return "answer_key"
    if "question_no" in canon and "response" in canon:
        return "results_long"
    if any(re.fullmatch(r"q\d+", c) for c in canon):
        return "results_responses"
    return "results_marks_only"


# ---------------------------------------------------------------- normalisers
GENDER = {"m": "male", "male": "male", "boy": "male", "f": "female", "female": "female", "girl": "female", "other": "other"}
EXAM = {"jee": "jee_main", "jee main": "jee_main", "jee_main": "jee_main", "jee-main": "jee_main", "neet": "neet_ug", "neet ug": "neet_ug", "neet_ug": "neet_ug", "neet-ug": "neet_ug"}
CLASS = {"11": "class_11", "xi": "class_11", "class 11": "class_11", "class_11": "class_11", "12": "class_12", "xii": "class_12", "class 12": "class_12",
         "class_12": "class_12", "dropper": "dropper", "droppers": "dropper", "repeater": "dropper"}
REL = {"father": "father", "mother": "mother", "pita": "father", "mata": "mother", "legal_guardian": "legal_guardian", "guardian": "legal_guardian", "self": "self"}
FEE_PLANS = {"upfront_full", "standard_3_instalment", "rajasthan_4_equal", "custom_scholarship"}


def norm(value: str, table: dict[str, str], default: str = "") -> str:
    v = str(value or "").strip().lower()
    return table.get(v, value if value in table.values() else default) if v else default


# ---------------------------------------------------------------- the parser
KIND_DEFAULT_TYPE = {"students": "students", "guardians": "guardians", "attendance": "attendance", "fees": "fees", "results": "results_marks_only", "answer-key": "answer_key"}


def _fallback_header(rows, override):
    """No row with 3 known headers: use the first row that has any text, so the caller can show a mapping screen."""
    for pos, (_r, vals, _h) in enumerate(rows):
        if any(v is not None and str(v).strip() for v in vals):
            colmap = {}
            for c_i, v in enumerate(vals):
                canon, conf, method = map_header(v, override)
                if canon:
                    colmap[c_i] = (canon, conf, method, str(v).strip())
            return pos, colmap
    raise ParseError("the file has no rows")


def parse(data: bytes, filename: str, override: dict[str, str] | None = None, max_marks_hint: float | None = None, kind: str | None = None) -> dict:
    rows = read_rows(data, filename)
    fallback = False
    try:
        hdr_pos, colmap = find_header(rows, override)
    except ParseError:
        if not kind:
            raise
        hdr_pos, colmap = _fallback_header(rows, override)
        fallback = True
    canon_set = {m[0] for m in colmap.values()}
    itype = detect_type(canon_set)
    if kind and (fallback or len(canon_set) < 3):
        itype = KIND_DEFAULT_TYPE[kind]
    qcols = sorted({c for c in canon_set if re.fullmatch(r"q\d+", c)}, key=lambda c: int(c[1:]))
    mapping = [{"column": m[3], "canonical": m[0], "confidence": m[1], "method": m[2]} for _, m in sorted(colmap.items())]
    header_cells = rows[hdr_pos][1]
    unmapped = [str(v).strip() for i, v in enumerate(header_cells) if i not in colmap and v is not None and str(v).strip()]
    missing = [c for c in REQUIRED[itype] if c not in canon_set]
    if itype == "results_marks_only" and not ({"raw_score", "physics_score"} & canon_set):
        missing.append("raw_score")
    out = {"ingest_type": itype, "mapping": mapping, "unmapped_columns": unmapped, "missing_required": missing,
           "question_columns": qcols, "canonical_rows": [], "errors": [], "header_row": rows[hdr_pos][0]}
    if missing:
        return out

    seen: set = set()
    for r_num, vals, hidden in rows[hdr_pos + 1:]:
        rd = {c: (vals[i] if i < len(vals) else None) for i, (c, *_rest) in colmap.items()}
        if hidden:
            out["errors"].append({"row_number": r_num, "code": "hidden_row_skipped", "field": "roll_no", "value": str(rd.get("roll_no") or "")})
            continue
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        first = str(vals[0] if vals else "").strip().lower()
        if any(t in first for t in ("grand total", "total attempted", "batch average", "class average")):
            out["errors"].append({"row_number": r_num, "code": "trailing_totals_row", "field": "test_name", "value": str(vals[0]).strip()})
            continue
        try:
            row, err = _row(itype, rd, qcols, seen, max_marks_hint)
        except (ValueError, TypeError) as ex:
            err = {"code": "unreadable_value", "field": "row", "value": str(ex)[:80]}
            row = None
        if err:
            out["errors"].append({"row_number": r_num, **err})
        if row is not None:
            out["canonical_rows"].append({"_row": r_num, **row})
    return out


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _row(itype: str, rd: dict, qcols: list[str], seen: set, max_marks_hint):
    if itype == "students":
        r_no = _s(rd.get("roll_no"))
        if not r_no:
            return None, {"code": "missing_roll_no", "field": "roll_no", "value": ""}
        if r_no in seen:
            return None, {"code": "duplicate_roll_no", "field": "roll_no", "value": r_no}
        dob, join = parse_date(rd.get("dob")), parse_date(rd.get("join_date"))
        if not dob:
            return None, {"code": "invalid_date", "field": "dob", "value": _s(rd.get("dob"))}
        if not join:
            return None, {"code": "invalid_date", "field": "join_date", "value": _s(rd.get("join_date"))}
        d_dob, d_join = dt.date.fromisoformat(dob), dt.date.fromisoformat(join)
        if d_join < dt.date(d_dob.year + 16, d_dob.month, min(28, d_dob.day)):
            return None, {"code": "underage_enrolment_lt_16", "field": "dob", "value": dob}
        seen.add(r_no)
        return {"roll_no": r_no, "name": _s(rd.get("name")), "name_alt": _s(rd.get("name_alt")), "dob": dob,
                "gender": norm(_s(rd.get("gender")), GENDER) or _s(rd.get("gender")),
                "class_level": norm(_s(rd.get("class_level")), CLASS) or _s(rd.get("class_level")),
                "exam": norm(_s(rd.get("exam")), EXAM) or _s(rd.get("exam")),
                "batch_name": _s(rd.get("batch_name")), "join_date": join, "fee_plan": _s(rd.get("fee_plan")),
                "out_of_town": bool(ref._parse_bool_flexible(rd.get("out_of_town"), ABS_SET, PRES_SET))}, None
    if itype == "attendance":
        r_no = _s(rd.get("roll_no"))
        if "MISSING" in r_no or "UNKNOWN" in r_no:
            return None, {"code": "missing_student_in_roster", "field": "roll_no", "value": r_no}
        d = parse_date(rd.get("attendance_date"))
        if d is None:
            return None, {"code": "invalid_date", "field": "attendance_date", "value": _s(rd.get("attendance_date"))}
        flag = ref._parse_bool_flexible(rd.get("present_flag"), ABS_SET, PRES_SET)
        if flag is None:
            return None, {"code": "invalid_attendance_marker", "field": "present_flag", "value": _s(rd.get("present_flag"))}
        return {"roll_no": r_no, "attendance_date": d, "present_flag": bool(flag), "source": _s(rd.get("source")) or "csv"}, None
    if itype == "results_marks_only":
        r_no = _s(rd.get("roll_no"))
        d = parse_date(rd.get("test_date"))
        absent = bool(ref._parse_bool_flexible(rd.get("absent_flag"), ABS_SET, PRES_SET))
        if _s(rd.get("physics_score")).lower() in ABS_SET or _s(rd.get("raw_score")).lower() in ABS_SET:
            absent = True
        base = {"test_name": _s(rd.get("test_name")), "test_date": d, "roll_no": r_no}
        if absent:
            return {**base, "raw_score": None, "physics_score": None, "chemistry_score": None, "math_or_bio_score": None, "absent_flag": True}, None
        p_s = _num(rd["physics_score"]) if _s(rd.get("physics_score")) else None
        c_s = _num(rd["chemistry_score"]) if _s(rd.get("chemistry_score")) else None
        m_s = _num(rd["math_or_bio_score"]) if _s(rd.get("math_or_bio_score")) else None
        raw_cell = _s(rd.get("raw_score"))
        if raw_cell.startswith("="):
            tot = sum(x for x in (p_s, c_s, m_s) if x is not None)
        elif raw_cell:
            tot = _num(raw_cell)
        else:
            tot = sum(x for x in (p_s, c_s, m_s) if x is not None)
        key = (base["test_name"], r_no)
        limit = max_marks_hint or 720.0
        if tot > limit or tot < -limit or key in seen:
            return None, {"code": "duplicate_or_out_of_bounds_score", "field": "raw_score", "value": str(tot)}
        seen.add(key)
        return {**base, "raw_score": tot, "physics_score": p_s, "chemistry_score": c_s, "math_or_bio_score": m_s, "absent_flag": False}, None
    if itype == "results_responses":
        r_no = _s(rd.get("roll_no"))
        if "UNKNOWN" in r_no or "MISSING" in r_no:
            return None, {"code": "missing_student_in_roster", "field": "roll_no", "value": r_no}
        out = {"test_name": _s(rd.get("test_name")), "test_date": parse_date(rd.get("test_date")), "roll_no": r_no}
        for qk in qcols:
            qv = _s(rd.get(qk))
            out[qk] = None if qv.lower() in BLANK_MARKERS else qv
        return out, None
    if itype == "results_long":
        qn = int(_num(rd.get("question_no")))
        qv = _s(rd.get("response"))
        return {"test_name": _s(rd.get("test_name")), "test_date": parse_date(rd.get("test_date")), "roll_no": _s(rd.get("roll_no")),
                "question_no": qn, "response": None if qv.lower() in BLANK_MARKERS else qv}, None
    if itype == "fees":
        r_no = _s(rd.get("roll_no"))
        if not r_no:
            return None, {"code": "missing_roll_no", "field": "roll_no", "value": ""}
        amt_str = _s(rd.get("amount_paise")).replace(",", "") or "0"
        amt = int(_num(amt_str))
        if amt < 0:
            return None, {"code": "negative_fee_amount", "field": "amount_paise", "value": amt_str}
        due = parse_date(rd.get("due_date"))
        if due is None:
            return None, {"code": "invalid_date", "field": "due_date", "value": _s(rd.get("due_date"))}
        paid_amt = int(_num(_s(rd.get("paid_amount_paise")) or "0"))
        if paid_amt < 0:
            return None, {"code": "negative_fee_amount", "field": "paid_amount_paise", "value": str(paid_amt)}
        return {"roll_no": r_no, "instalment_no": int(_num(rd.get("instalment_no"))), "due_date": due, "amount_paise": amt,
                "paid_date": parse_date(rd.get("paid_date")), "paid_amount_paise": paid_amt,
                "refund_due_paise": int(_num(_s(rd.get("refund_due_paise")) or "0")),
                "refund_paid_date": parse_date(rd.get("refund_paid_date"))}, None
    if itype == "guardians":
        ph = re.sub(r"[\s\-+]", "", _s(rd.get("phone")))
        ph = ph[-10:] if len(ph) in (11, 12, 13) and ph[-10:].isdigit() else ph
        if not re.match(r"^[0-9]{10}$", ph):
            return None, {"code": "invalid_phone_number", "field": "phone", "value": _s(rd.get("phone"))}
        return {"roll_no": _s(rd.get("roll_no")), "guardian_name": _s(rd.get("guardian_name")), "phone": ph,
                "relation": norm(_s(rd.get("relation")), REL) or _s(rd.get("relation")),
                "lang_pref": (_s(rd.get("lang_pref")).lower()[:2] or "hi")}, None
    if itype == "answer_key":
        q_no = int(_num(rd.get("question_no")))
        ans_raw = _s(rd.get("accepted_answers"))
        ans = [x.strip() for x in ans_raw.split("|") if x.strip()]
        q_t = _s(rd.get("q_type")) or ("mcq_single" if ans and all(re.fullmatch(r"[A-Da-d]", a) for a in ans) and len(ans) == 1 else "numeric")
        if q_t == "numeric":
            for a in ans:
                float(a)
        if q_t in ("mcq_single", "mcq_multi") and (not ans or any(x not in ("A", "B", "C", "D") for x in ans)):
            return None, {"code": "invalid_option_label", "field": "accepted_answers", "value": ans_raw}
        if q_no in seen:
            return None, {"code": "duplicate_question_no", "field": "question_no", "value": str(q_no)}
        seen.add(q_no)
        return {"question_no": q_no, "q_type": q_t, "subject": _s(rd.get("subject")), "topic_id": _s(rd.get("topic_id")),
                "accepted_answers": ans, "numeric_tolerance": _num(_s(rd.get("numeric_tolerance")) or "0.0"),
                "dropped": bool(ref._parse_bool_flexible(rd.get("dropped"), ABS_SET, PRES_SET))}, None
    raise ParseError(f"unhandled type {itype}")
